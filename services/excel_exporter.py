from io import BytesIO
import json
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def generate_results_excel(test_title: str, submissions: List[Dict[str, Any]]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"

    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Title row
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"Ведомость результатов: {test_title}"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "№", "ФИО Ученика", "Класс", "Баллы", "Всего вопросов", "Результат (%)", "Дата сдачи"
    ]
    ws.append([]) # row 2 empty
    ws.append(headers) # row 3
    ws.row_dimensions[3].height = 24

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for idx, sub in enumerate(submissions, start=1):
        score = sub.get("score", 0)
        total = sub.get("total_questions", 1)
        pct = round((score / total) * 100, 1) if total > 0 else 0
        date_str = str(sub.get("completed_at", ""))[:16]

        row = [
            idx,
            sub.get("full_name", "Не указано"),
            sub.get("class_name", "Не указан"),
            score,
            total,
            f"{pct}%",
            date_str
        ]
        ws.append(row)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        # formatting cells
        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_num)
            c.border = thin_border
            if col_num in (1, 3, 4, 5, 6, 7):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Fixed column widths
    col_widths = {1: 8, 2: 28, 3: 14, 4: 12, 5: 16, 6: 16, 7: 20}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
