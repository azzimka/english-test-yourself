import logging
from io import BytesIO
from typing import Tuple, Optional
import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)

class UnsupportedFileError(Exception):
    """Raised when file format is not supported."""
    pass

class FileExtractionError(Exception):
    """Raised when text cannot be extracted from a file."""
    pass

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from all pages of a PDF document."""
    try:
        reader = PdfReader(BytesIO(file_bytes))
        pages_text = []
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text()
            if text and text.strip():
                pages_text.append(text.strip())
        
        full_text = "\n\n".join(pages_text).strip()
        if not full_text:
            raise FileExtractionError("Не удалось извлечь текст из PDF (возможно, это сканированные картинки без текстового слоя).")
        return full_text
    except FileExtractionError:
        raise
    except Exception as e:
        logger.warning("Error reading PDF: %s", e)
        raise FileExtractionError(f"Ошибка при чтении PDF-документа: {e}")

def extract_text_from_docx_bytes(file_bytes: bytes) -> str:
    """Extracts text from paragraphs and tables of a docx file."""
    try:
        doc = docx.Document(BytesIO(file_bytes))
        full_text = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                full_text.append(text)

        # Extract text from tables if present
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        result = "\n".join(full_text).strip()
        if not result:
            raise FileExtractionError("Документ Word не содержит текста.")
        return result
    except FileExtractionError:
        raise
    except Exception as e:
        logger.warning("Error reading DOCX: %s", e)
        raise FileExtractionError(f"Ошибка при чтении Word (.docx) документа: {e}")

def extract_text_from_plain_text(file_bytes: bytes) -> str:
    """Extracts text with automatic encoding detection (utf-8, cp1251, latin1)."""
    encodings = ["utf-8", "cp1251", "windows-1251", "latin1"]
    for enc in encodings:
        try:
            text = file_bytes.decode(enc).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise FileExtractionError("Не удалось декодировать текстовый файл.")

def extract_text_from_file(file_name: str, file_bytes: bytes) -> str:
    """
    Main dispatcher for extracting text from supported document types.
    Raises UnsupportedFileError or FileExtractionError.
    """
    lower_name = (file_name or "").lower()

    if lower_name.endswith(".docx"):
        return extract_text_from_docx_bytes(file_bytes)
    elif lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith((".txt", ".md", ".rtf", ".csv")):
        return extract_text_from_plain_text(file_bytes)
    elif lower_name.endswith(".doc"):
        raise UnsupportedFileError(
            "Файл в устаревшем бинарном формате <b>.doc</b> (Word 97-2003).\n\n"
            "Пожалуйста, откройте его в Word и сохраните в современном формате <b>.docx</b> или <b>.pdf</b>, затем отправьте снова."
        )
    elif lower_name.endswith((".jpg", ".jpeg", ".png", ".webp", ".heic")):
        raise UnsupportedFileError(
            "📷 <b>Анализ изображений не поддерживается</b>\n\n"
            "Пожалуйста, отправьте документ в формате <b>PDF</b>, <b>Word (.docx)</b> или просто скопируйте текст задания в сообщение."
        )
    else:
        raise UnsupportedFileError(
            f"❌ <b>Формат файла не поддерживается</b>\n\n"
            "Бот умеет анализировать следующие типы материалов:\n"
            "• Документы <b>PDF (.pdf)</b>\n"
            "• Документы <b>Word (.docx)</b>\n"
            "• Текстовые файлы <b>(.txt)</b>\n"
            "• Текст сообщением в чат"
        )
