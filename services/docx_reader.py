import docx
from io import BytesIO

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from paragraphs and tables of a docx file."""
    doc = docx.Document(BytesIO(file_bytes))
    full_text = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_text.append(text)

    # Also extract text from any tables if present (often used for tests)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))

    return "\n".join(full_text)
